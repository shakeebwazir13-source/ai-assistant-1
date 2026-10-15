 # ATS Resume Analyzer - Streamlit + Google Gemini Flash
# Upload a resume (PDF / DOCX / TXT) and get an ATS score and improvement tips.
import json
import os
import re
from io import BytesIO

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
# Primary model first; if Google retires it (404), the next one is tried automatically.
MODEL_CANDIDATES = [
    m for m in (os.getenv("GEMINI_MODEL"), "gemini-3.8-flash", "gemini-flash-latest") if m
]
MAX_RESUME_CHARS = 15000  # keeps the prompt small and fast
MAX_FILE_MB = 5
MIN_TEXT_CHARS = 200  # below this we assume the file has no readable text

SCORE_KEYS = {
    "formatting": "Formatting & Parsability",
    "keywords": "Keywords",
    "content_quality": "Content Quality",
    "structure": "Structure & Sections",
    "impact": "Impact & Achievements",
}


# ----------------------------------------------------------------------------
# File reading
# ----------------------------------------------------------------------------
def extract_text(file_name: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX or TXT file."""
    name = file_name.lower()

    if name.endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()

    if name.endswith(".docx"):
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# ----------------------------------------------------------------------------
# Gemini
# ----------------------------------------------------------------------------
def build_prompt(resume_text: str, job_description: str) -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n{job_description.strip()[:6000]}\n"
        if job_description.strip()
        else "JOB DESCRIPTION: (not provided - evaluate for general ATS-friendliness "
        "and for common keywords in the candidate's apparent field)\n"
    )
    return f"""You are an expert ATS (Applicant Tracking System) and resume reviewer.
Analyze the resume below and return ONLY a valid JSON object (no markdown, no commentary).

Scoring rules:
- All scores are integers from 0 to 100. Be realistic and strict; most resumes score 50-80.
- "overall_score" is the overall ATS score. If a job description is given, weigh keyword match to it heavily.
- Only list items in "keywords_found" that really appear in the resume.
- Never invent experience, degrees or skills the candidate does not have.

JSON schema:
{{
  "overall_score": 0,
  "scores": {{
    "formatting": 0,
    "keywords": 0,
    "content_quality": 0,
    "structure": 0,
    "impact": 0
  }},
  "summary": "2-3 sentence overall assessment",
  "strengths": ["..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "area": "section or topic", "issue": "what is wrong", "suggestion": "specific fix"}}
  ],
  "keywords_found": ["..."],
  "keywords_missing": ["..."],
  "rewrite_examples": [
    {{"before": "weak bullet copied from the resume", "after": "stronger version using action verb + metric (use placeholders like [X]% if numbers are unknown)"}}
  ],
  "formatting_tips": ["..."]
}}

Give 5-8 improvements ordered by priority, up to 15 keywords_missing, and 2-4 rewrite_examples.

{jd_block}
RESUME:
{resume_text[:MAX_RESUME_CHARS]}
"""


def parse_json(text: str) -> dict:
    """Parse model output into a dict, tolerating markdown fences / extra text."""
    text = (text or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start != -1 and end > start:
            return json.loads(text[start : end + 1])
        raise


def _score(value) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def _str_list(value) -> list:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def normalize_result(data: dict) -> dict:
    """Make sure every field exists and has the right type so the UI never crashes."""
    if not isinstance(data, dict):
        data = {}
    raw_scores = data.get("scores") if isinstance(data.get("scores"), dict) else {}
    scores = {k: _score(raw_scores.get(k)) for k in SCORE_KEYS}

    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            priority = str(item.get("priority", "Medium")).strip().capitalize()
            if priority not in ("High", "Medium", "Low"):
                priority = "Medium"
            improvements.append(
                {
                    "priority": priority,
                    "area": str(item.get("area", "General")).strip(),
                    "issue": str(item.get("issue", "")).strip(),
                    "suggestion": str(item.get("suggestion", "")).strip(),
                }
            )
        elif isinstance(item, str) and item.strip():
            improvements.append(
                {"priority": "Medium", "area": "General", "issue": "", "suggestion": item.strip()}
            )

    examples = []
    for item in data.get("rewrite_examples") or []:
        if isinstance(item, dict) and (item.get("before") or item.get("after")):
            examples.append(
                {"before": str(item.get("before", "")).strip(), "after": str(item.get("after", "")).strip()}
            )

    overall = data.get("overall_score")
    if overall is None and scores:
        overall = sum(scores.values()) / len(scores)

    return {
        "overall_score": _score(overall),
        "scores": scores,
        "summary": str(data.get("summary", "")).strip(),
        "strengths": _str_list(data.get("strengths")),
        "improvements": improvements,
        "keywords_found": _str_list(data.get("keywords_found")),
        "keywords_missing": _str_list(data.get("keywords_missing")),
        "rewrite_examples": examples,
        "formatting_tips": _str_list(data.get("formatting_tips")),
    }


def analyze_resume(api_key: str, resume_text: str, job_description: str) -> dict:
    client = genai.Client(api_key=api_key)
    prompt = build_prompt(resume_text, job_description)
    config = types.GenerateContentConfig(
        temperature=0.2,
        response_mime_type="application/json",
    )
    last_error = None
    for model in MODEL_CANDIDATES:
        try:
            response = client.models.generate_content(model=model, contents=prompt, config=config)
            return normalize_result(parse_json(response.text))
        except Exception as e:  # try the next model only if this one is unavailable
            msg = str(e)
            if "404" in msg or "NOT_FOUND" in msg:
                last_error = e
                continue
            raise
    raise last_error


# ----------------------------------------------------------------------------
# UI helpers
# ----------------------------------------------------------------------------
def get_api_key() -> str:
    """Read the key from Streamlit secrets or an environment variable."""
    try:
        key = st.secrets.get("GEMINI_API_KEY")
    except Exception:
        key = None
    return key or os.getenv("GEMINI_API_KEY") or ""


def score_label(score: int) -> str:
    if score >= 80:
        return "🟢 Excellent"
    if score >= 60:
        return "🟡 Good - needs work"
    return "🔴 Needs major improvement"


def render_results(result: dict) -> None:
    overall = result["overall_score"]

    st.subheader("Your ATS Score")
    col1, col2 = st.columns([1, 2])
    col1.metric("Overall", f"{overall}/100")
    col2.markdown(f"**{score_label(overall)}**")
    col2.progress(overall / 100)
    if result["summary"]:
        st.write(result["summary"])

    st.subheader("Score Breakdown")
    for key, label in SCORE_KEYS.items():
        value = result["scores"][key]
        st.write(f"{label}: **{value}/100**")
        st.progress(value / 100)

    if result["strengths"]:
        st.subheader("✅ Strengths")
        for s in result["strengths"]:
            st.markdown(f"- {s}")

    if result["improvements"]:
        st.subheader("🛠️ Improvements")
        order = {"High": 0, "Medium": 1, "Low": 2}
        icons = {"High": "🔴", "Medium": "🟡", "Low": "🟢"}
        for item in sorted(result["improvements"], key=lambda i: order[i["priority"]]):
            with st.expander(f"{icons[item['priority']]} {item['priority']} - {item['area']}"):
                if item["issue"]:
                    st.markdown(f"**Issue:** {item['issue']}")
                if item["suggestion"]:
                    st.markdown(f"**Fix:** {item['suggestion']}")

    st.subheader("🔑 Keywords")
    kcol1, kcol2 = st.columns(2)
    with kcol1:
        st.markdown("**Found**")
        st.write(", ".join(result["keywords_found"]) or "None detected")
    with kcol2:
        st.markdown("**Missing / recommended**")
        st.write(", ".join(result["keywords_missing"]) or "None")

    if result["rewrite_examples"]:
        st.subheader("✍️ Example Rewrites")
        for ex in result["rewrite_examples"]:
            st.markdown(f"**Before:** {ex['before']}")
            st.markdown(f"**After:** {ex['after']}")
            st.divider()

    if result["formatting_tips"]:
        st.subheader("📄 Formatting Tips")
        for tip in result["formatting_tips"]:
            st.markdown(f"- {tip}")


# ----------------------------------------------------------------------------
# App
# ----------------------------------------------------------------------------
def main() -> None:
    st.set_page_config(page_title="ATS Resume Analyzer", page_icon="📄", layout="centered")
    st.title("📄 ATS Resume Analyzer")
    st.caption("Upload your resume to get an ATS score and tips to improve it. Powered by Google Gemini.")

    api_key = get_api_key()
    if not api_key:
        with st.sidebar:
            st.header("Settings")
            api_key = st.text_input(
                "Gemini API key",
                type="password",
                help="Get a free key at https://aistudio.google.com/apikey",
            )
        if not api_key:
            st.info("Enter your Gemini API key in the sidebar to begin.")

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional)",
        height=150,
        placeholder="Paste the job description here for a tailored keyword match...",
    )

    if st.button("Analyze Resume", type="primary", disabled=not (uploaded and api_key)):
        if uploaded.size > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
            return

        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded.name, uploaded.getvalue())
        except Exception as e:
            st.error(f"Could not read the file: {e}")
            return

        if len(text) < MIN_TEXT_CHARS:
            st.error(
                "Could not find enough text in this file. If it is a scanned image PDF, "
                "export a text-based PDF or upload a DOCX instead."
            )
            return

        try:
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(api_key, text, job_description)
        except json.JSONDecodeError:
            st.error("The AI returned an unreadable response. Please try again.")
            return
        except Exception as e:
            st.error(f"Gemini request failed: {e}")
            return

        render_results(result)
        st.download_button(
            "Download report (JSON)",
            data=json.dumps(result, indent=2),
            file_name="ats_report.json",
            mime="application/json",
        )


if __name__ == "__main__":
    main()
